from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from .local_store import LocalKnowledgeStore

SUPPORTED_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".log",
    ".csv",
    ".xlsx",
    ".xlsm",
    ".docx",
    ".pdf",
    ".pptx",
}
TEXT_KEYS = {
    "text",
    "content",
    "message",
    "prompt",
    "response",
    "question",
    "answer",
    "output",
    "input",
}


@dataclass
class HistoryImportStats:
    scanned_files: int = 0
    imported_files: int = 0
    skipped_files: int = 0
    extracted_facts: int = 0
    duplicate_skipped: int = 0
    failed_files: list[str] | None = None
    source_counts: dict[str, int] | None = None

    def to_dict(self) -> dict:
        return {
            "scanned_files": self.scanned_files,
            "imported_files": self.imported_files,
            "skipped_files": self.skipped_files,
            "extracted_facts": self.extracted_facts,
            "duplicate_skipped": self.duplicate_skipped,
            "failed_files": self.failed_files or [],
            "source_counts": self.source_counts or {},
        }


def ingest_history_directory(
    store: LocalKnowledgeStore,
    directory: Path,
    *,
    recursive: bool = True,
    max_files: int = 500,
    max_lines_per_file: int = 80,
    max_chars_per_fact: int = 300,
    progress_callback: Callable[[dict], None] | None = None,
    source_files: list[Path] | None = None,
) -> HistoryImportStats:
    stats = HistoryImportStats(source_counts={}, failed_files=[])
    facts: list[str] = []
    seen_in_run: set[str] = set()
    existing_facts = set(store.load().doc_facts)

    files = source_files or _scan_files(directory, recursive=recursive, max_files=max_files)
    stats.scanned_files = len(files)
    total = len(files)
    for idx, path in enumerate(files, start=1):
        source = detect_history_source(path)
        stats.source_counts[source] = stats.source_counts.get(source, 0) + 1
        try:
            extracted = extract_facts_from_history_file(
                path,
                max_lines=max_lines_per_file,
                max_chars_per_fact=max_chars_per_fact,
            )
        except OSError:
            stats.failed_files.append(str(path))
            extracted = []
        if not extracted:
            stats.skipped_files += 1
        else:
            stats.imported_files += 1
            for item in extracted:
                if item in seen_in_run:
                    stats.duplicate_skipped += 1
                    continue
                seen_in_run.add(item)
                if item in existing_facts:
                    stats.duplicate_skipped += 1
                    continue
                facts.append(item)
        if progress_callback:
            progress_callback(
                {
                    "current": idx,
                    "total": total,
                    "path": str(path),
                    "imported_files": stats.imported_files,
                    "skipped_files": stats.skipped_files,
                    "facts": len(facts),
                    "failed_files": len(stats.failed_files),
                }
            )

    if facts:
        store.append_doc_facts(facts)
    stats.extracted_facts = len(facts)
    return stats


def scan_history_files(directory: Path, *, recursive: bool = True, max_files: int = 500) -> list[Path]:
    return _scan_files(directory, recursive=recursive, max_files=max_files)


def _scan_files(directory: Path, *, recursive: bool, max_files: int) -> list[Path]:
    pattern = "**/*" if recursive else "*"
    out: list[Path] = []
    for path in directory.glob(pattern):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        out.append(path)
        if len(out) >= max_files:
            break
    return out


def detect_history_source(path: Path) -> str:
    s = str(path).lower()
    if "gemini" in s:
        return "gemini"
    if "deepseek" in s:
        return "deepseek"
    if "doubao" in s:
        return "doubao"
    if "qwen" in s or "aliyun" in s:
        return "qwen"
    if "chatgpt" in s or "openai" in s:
        return "chatgpt"
    if "claude" in s:
        return "claude"
    if "kimi" in s:
        return "kimi"
    if "coze" in s or "扣子" in s:
        return "coze"
    return "unknown"


def extract_facts_from_history_file(path: Path, *, max_lines: int, max_chars_per_fact: int) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        return _extract_from_json(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_from_excel(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".docx":
        return _extract_from_docx(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".pdf":
        return _extract_from_pdf(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".pptx":
        return _extract_from_pptx(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    return _extract_from_plain_text(path, max_lines=max_lines, max_chars=max_chars_per_fact)


def _extract_from_plain_text(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    lines = [_clean_line(line, max_chars=max_chars) for line in raw.splitlines()]
    return [line for line in lines if line][:max_lines]


def _extract_from_json(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return _extract_from_plain_text(path, max_lines=max_lines, max_chars=max_chars)
    collected: list[str] = []
    for value in _walk_json_strings(data):
        line = _clean_line(value, max_chars=max_chars)
        if line:
            collected.append(line)
        if len(collected) >= max_lines:
            break
    return collected


def _extract_from_excel(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    try:
        wb = load_workbook(filename=path, data_only=True, read_only=True)
    except Exception:
        return []
    collected: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if not cells:
                    continue
                line = _clean_line(" | ".join(cells), max_chars=max_chars)
                if line:
                    collected.append(line)
                if len(collected) >= max_lines:
                    return collected
    finally:
        wb.close()
    return collected


def _extract_from_docx(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    try:
        doc = Document(path)
    except Exception:
        return []
    collected: list[str] = []
    for paragraph in doc.paragraphs:
        line = _clean_line(paragraph.text, max_chars=max_chars)
        if not line:
            continue
        collected.append(line)
        if len(collected) >= max_lines:
            break
    return collected


def _extract_from_pdf(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    try:
        reader = PdfReader(str(path))
    except Exception:
        return []
    collected: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for raw_line in text.splitlines():
            line = _clean_line(raw_line, max_chars=max_chars)
            if not line:
                continue
            collected.append(line)
            if len(collected) >= max_lines:
                return collected
    return collected


def _extract_from_pptx(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    try:
        prs = Presentation(str(path))
    except Exception:
        return []
    collected: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            line = _clean_line(text, max_chars=max_chars)
            if not line:
                continue
            collected.append(line)
            if len(collected) >= max_lines:
                return collected
    return collected


def _walk_json_strings(value: object) -> Iterable[str]:
    if isinstance(value, str):
        yield value
        return
    if isinstance(value, list):
        for item in value:
            yield from _walk_json_strings(item)
        return
    if isinstance(value, dict):
        for key, item in value.items():
            if isinstance(item, str):
                if key.lower() in TEXT_KEYS:
                    yield item
                elif len(item) > 24:
                    yield item
            else:
                yield from _walk_json_strings(item)


def _clean_line(line: str, *, max_chars: int) -> str:
    text = line.strip()
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    if text.startswith(("{", "}", "[", "]")):
        return ""
    return text[:max_chars]

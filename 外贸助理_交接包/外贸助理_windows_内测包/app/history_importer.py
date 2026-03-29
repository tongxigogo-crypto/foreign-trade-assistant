from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

from .local_store import LocalKnowledgeStore

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None


SUPPORTED_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".log",
    ".csv",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
}
TEXT_KEYS = {
    "text",
    "content",
    "message",
    "prompt",
    "response",
    "question",
    "answer",
    "output",
    "input",
}


@dataclass
class HistoryImportStats:
    scanned_files: int = 0
    imported_files: int = 0
    skipped_files: int = 0
    extracted_facts: int = 0
    duplicate_skipped: int = 0
    failed_files: list[str] | None = None
    source_counts: dict[str, int] | None = None

    def to_dict(self) -> dict:
        return {
            "scanned_files": self.scanned_files,
            "imported_files": self.imported_files,
            "skipped_files": self.skipped_files,
            "extracted_facts": self.extracted_facts,
            "duplicate_skipped": self.duplicate_skipped,
            "failed_files": self.failed_files or [],
            "source_counts": self.source_counts or {},
        }


def ingest_history_directory(
    store: LocalKnowledgeStore,
    directory: Path,
    *,
    recursive: bool = True,
    max_files: int = 500,
    max_lines_per_file: int = 80,
    max_chars_per_fact: int = 300,
    progress_callback: Callable[[dict], None] | None = None,
    source_files: list[Path] | None = None,
) -> HistoryImportStats:
    stats = HistoryImportStats(source_counts={}, failed_files=[])
    facts: list[str] = []
    seen_in_run: set[str] = set()
    existing_facts = set(store.load().doc_facts)

    files = source_files or _scan_files(directory, recursive=recursive, max_files=max_files)
    stats.scanned_files = len(files)
    total = len(files)
    for idx, path in enumerate(files, start=1):
        source = detect_history_source(path)
        stats.source_counts[source] = stats.source_counts.get(source, 0) + 1
        try:
            extracted = extract_facts_from_history_file(
                path,
                max_lines=max_lines_per_file,
                max_chars_per_fact=max_chars_per_fact,
            )
        except Exception:
            stats.failed_files.append(str(path))
            extracted = []
        if not extracted:
            stats.skipped_files += 1
        else:
            stats.imported_files += 1
            for item in extracted:
                if item in seen_in_run:
                    stats.duplicate_skipped += 1
                    continue
                seen_in_run.add(item)
                if item in existing_facts:
                    stats.duplicate_skipped += 1
                    continue
                facts.append(item)
        if progress_callback:
            progress_callback(
                {
                    "current": idx,
                    "total": total,
                    "path": str(path),
                    "imported_files": stats.imported_files,
                    "skipped_files": stats.skipped_files,
                    "facts": len(facts),
                    "failed_files": len(stats.failed_files),
                }
            )

    if facts:
        store.append_doc_facts(facts)
    stats.extracted_facts = len(facts)
    return stats


def scan_history_files(directory: Path, *, recursive: bool = True, max_files: int = 500) -> list[Path]:
    return _scan_files(directory, recursive=recursive, max_files=max_files)


def _scan_files(directory: Path, *, recursive: bool, max_files: int) -> list[Path]:
    pattern = "**/*" if recursive else "*"
    out: list[Path] = []
    for path in directory.glob(pattern):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        out.append(path)
        if len(out) >= max_files:
            break
    return out


def detect_history_source(path: Path) -> str:
    s = str(path).lower()
    if "gemini" in s:
        return "gemini"
    if "deepseek" in s:
        return "deepseek"
    if "doubao" in s:
        return "doubao"
    if "qwen" in s or "aliyun" in s:
        return "qwen"
    if "chatgpt" in s or "openai" in s:
        return "chatgpt"
    if "claude" in s:
        return "claude"
    if "kimi" in s:
        return "kimi"
    if "coze" in s or "扣子" in s:
        return "coze"
    return "unknown"


def extract_facts_from_history_file(path: Path, *, max_lines: int, max_chars_per_fact: int) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        return _extract_from_json(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".pdf":
        return _extract_from_pdf(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".docx":
        return _extract_from_docx(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".xlsx":
        return _extract_from_xlsx(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    if suffix == ".xls":
        return _extract_from_xls(path, max_lines=max_lines, max_chars=max_chars_per_fact)
    return _extract_from_plain_text(path, max_lines=max_lines, max_chars=max_chars_per_fact)


def _extract_from_plain_text(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    lines = [_clean_line(line, max_chars=max_chars) for line in raw.splitlines()]
    return [line for line in lines if line][:max_lines]


def _extract_from_json(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return _extract_from_plain_text(path, max_lines=max_lines, max_chars=max_chars)
    collected: list[str] = []
    for value in _walk_json_strings(data):
        line = _clean_line(value, max_chars=max_chars)
        if line:
            collected.append(line)
        if len(collected) >= max_lines:
            break
    return collected


def _extract_from_pdf(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    if PdfReader is None:
        return []
    try:
        reader = PdfReader(path)
    except Exception:
        return []
    collected: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            continue
        for line in text.splitlines():
            cleaned = _clean_line(line, max_chars=max_chars)
            if cleaned:
                collected.append(cleaned)
            if len(collected) >= max_lines:
                return collected
    return collected


def _extract_from_docx(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    if Document is None:
        return []
    try:
        document = Document(path)
    except Exception:
        return []
    collected: list[str] = []
    for paragraph in document.paragraphs:
        cleaned = _clean_line(paragraph.text, max_chars=max_chars)
        if cleaned:
            collected.append(cleaned)
        if len(collected) >= max_lines:
            break
    return collected


def _extract_from_xlsx(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    if load_workbook is None:
        return []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        return _extract_rows_from_sheets(
            (sheet.iter_rows(values_only=True) for sheet in workbook.worksheets),
            max_lines=max_lines,
            max_chars=max_chars,
        )
    finally:
        close = getattr(workbook, "close", None)
        if callable(close):
            close()


def _extract_from_xls(path: Path, *, max_lines: int, max_chars: int) -> list[str]:
    if xlrd is None:
        return []
    try:
        workbook = xlrd.open_workbook(path)
    except Exception:
        return []
    sheets = []
    for index in range(getattr(workbook, "nsheets", 0)):
        try:
            sheet = workbook.sheet_by_index(index)
        except Exception:
            continue
        sheets.append((sheet.row_values(row_index) for row_index in range(getattr(sheet, "nrows", 0))))
    return _extract_rows_from_sheets(sheets, max_lines=max_lines, max_chars=max_chars)


def _extract_rows_from_sheets(sheet_rows: Iterable[Iterable[object]], *, max_lines: int, max_chars: int) -> list[str]:
    collected: list[str] = []
    for rows in sheet_rows:
        for row in rows:
            line = " ".join(str(cell).strip() for cell in row if str(cell).strip())
            cleaned = _clean_line(line, max_chars=max_chars)
            if cleaned:
                collected.append(cleaned)
            if len(collected) >= max_lines:
                return collected
    return collected


def _walk_json_strings(value: object) -> Iterable[str]:
    if isinstance(value, str):
        yield value
        return
    if isinstance(value, list):
        for item in value:
            yield from _walk_json_strings(item)
        return
    if isinstance(value, dict):
        for key, item in value.items():
            if isinstance(item, str):
                if key.lower() in TEXT_KEYS:
                    yield item
                elif len(item) > 24:
                    yield item
            else:
                yield from _walk_json_strings(item)


def _clean_line(line: str, *, max_chars: int) -> str:
    text = line.strip()
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    if text.startswith(("{", "}", "[", "]")):
        return ""
    return text[:max_chars]

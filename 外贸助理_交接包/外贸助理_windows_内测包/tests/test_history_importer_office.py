from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pptx import Presentation

from app.history_importer import extract_facts_from_history_file, scan_history_files


def test_scan_history_files_includes_office_suffixes(tmp_path: Path) -> None:
    (tmp_path / "sample.xlsx").write_text("", encoding="utf-8")
    (tmp_path / "sample.docx").write_text("", encoding="utf-8")
    (tmp_path / "sample.pptx").write_text("", encoding="utf-8")
    (tmp_path / "sample.pdf").write_text("", encoding="utf-8")
    (tmp_path / "sample.doc").write_text("", encoding="utf-8")
    (tmp_path / "sample.ppt").write_text("", encoding="utf-8")
    (tmp_path / "sample.xls").write_text("", encoding="utf-8")
    (tmp_path / "ignore.bin").write_text("", encoding="utf-8")

    files = scan_history_files(tmp_path, recursive=False, max_files=50)
    names = sorted(path.name for path in files)

    assert "sample.xlsx" in names
    assert "sample.docx" in names
    assert "sample.pptx" in names
    assert "sample.pdf" in names
    assert "sample.doc" in names
    assert "sample.ppt" in names
    assert "sample.xls" in names
    assert "ignore.bin" not in names


def test_extract_facts_from_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "客户"
    ws["B1"] = "Ali Trading"
    ws["A2"] = "产品"
    ws["B2"] = "A100"
    wb.save(path)
    wb.close()

    facts = extract_facts_from_history_file(path, max_lines=20, max_chars_per_fact=200)
    joined = " ".join(facts)
    assert "Ali Trading" in joined
    assert "A100" in joined


def test_extract_facts_from_docx(tmp_path: Path) -> None:
    path = tmp_path / "memo.docx"
    doc = Document()
    doc.add_paragraph("本周报价需要保留 8% 毛利。")
    doc.add_paragraph("客户: Blue Ocean")
    doc.save(path)

    facts = extract_facts_from_history_file(path, max_lines=20, max_chars_per_fact=200)
    joined = " ".join(facts)
    assert "Blue Ocean" in joined


def test_extract_facts_from_pptx(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Follow-up"
    slide.placeholders[1].text = "客户反馈：需要更快交期。"
    prs.save(path)

    facts = extract_facts_from_history_file(path, max_lines=20, max_chars_per_fact=200)
    joined = " ".join(facts)
    assert "Quarterly Follow-up" in joined


def test_extract_facts_from_pdf_handles_file(tmp_path: Path) -> None:
    path = tmp_path / "note.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=200)
    with path.open("wb") as f:
        writer.write(f)

    facts = extract_facts_from_history_file(path, max_lines=20, max_chars_per_fact=200)
    assert isinstance(facts, list)


def test_extract_facts_from_legacy_doc_fallback(tmp_path: Path) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\x00\x00legacy-doc")
    facts = extract_facts_from_history_file(path, max_lines=20, max_chars_per_fact=200)
    assert facts
    assert "Legacy Office file detected" in facts[0]
